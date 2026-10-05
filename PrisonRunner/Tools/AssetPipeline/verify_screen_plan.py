"""Check final Word structure, rendered pages, and local gallery links."""
from pathlib import Path
import re
import argparse
import zipfile
from lxml import etree
from docx import Document
from pypdf import PdfReader

root = Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser()
parser.add_argument('--version',default='V10',choices=['V10','V11','V12','V13'])
args=parser.parse_args()
document = Document(root / f'docs/무한옥_화면별_게임기획서_{args.version}.docx')
assert len(document.inline_shapes) == (30 if args.version in ('V12','V13') else 20)
assert len(document.tables) == (4 if args.version in ('V11','V12','V13') else 1)
assert len(document.tables[0].rows) == 9
review = root / f'docs/document-qa/{args.version}'
if args.version in ('V11','V12','V13'):review=review/'final'
pdf = PdfReader(review / f'무한옥_화면별_게임기획서_{args.version}.pdf')
expected={'V10':21,'V11':24,'V12':29,'V13':29}[args.version]
assert len(pdf.pages) == expected, len(pdf.pages)
for number, page in enumerate(pdf.pages, 1):
    page_text = page.extract_text()
    assert len(page_text) > 200, number
    if args.version == 'V11':
        assert re.search(rf'무한옥\s+V11\s+·\s+{number}\b', page_text), (number, 'footer')
    assert (review / f'page-{number}.png').exists(), number
if args.version in ('V11','V12'):
    assert len(document.tables[1].rows)==11
    assert len(document.tables[2].rows)==6
    assert len(document.tables[3].rows)==6
    assert '22 맵 순서와 운동 배치' in pdf.pages[21].extract_text()
    assert '23 진행 거리에 따른 속도와 안내' in pdf.pages[22].extract_text()
    assert '24 조작 기준과 구현 검수' in pdf.pages[23].extract_text()
    body='\n'.join(p.text for p in document.paragraphs)
    for phrase in ['960 m','17 m/s','9 m/s','120 m','23 m','점프가 우선','V10 촬영본']:
        assert phrase in body, phrase
if args.version in ('V12','V13'):
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(root / f'docs/무한옥_화면별_게임기획서_{args.version}.docx') as package:
        body_xml=etree.fromstring(package.read('word/document.xml'))
        assert not body_xml.xpath('//w:footerReference',namespaces=ns)
        for name in package.namelist():
            if name.startswith('word/footer') and name.endswith('.xml'):
                xml=etree.fromstring(package.read(name))
                assert not xml.xpath('//w:t|//w:fldSimple|//w:instrText|//w:drawing|//w:tbl',namespaces=ns),name
        source_version='V12' if args.version=='V13' else 'V11'
        with zipfile.ZipFile(root / f'docs/무한옥_화면별_게임기획서_{source_version}.docx') as source:
            for name in source.namelist():
                if name.startswith('word/media/'):
                    assert package.read(name)==source.read(name),name
    for page in pdf.pages:
        assert not re.search(r'무한옥\s+V1[123]\s+·',page.extract_text())
    for n,title in enumerate(['점프와 파이프 슬라이드','하이니와 입구 문 슬라이드','문 회피와 수레 균형','바닥 아이템 피하기','보조 장애물의 회피 기준'],25):
        assert f'{n} {title}' in pdf.pages[n-1].extract_text(),n
    for kind in ['crate','pipe','stairs','slidegate','door','rail','banana','puddle','barrier','laser']:
        assert (root/f'docs/document-qa/V12/obstacle-diagrams/{kind}.png').exists(),kind
if args.version=='V13':
    for table, count in zip(document.tables,[9,11,6,6]):
        assert len(table.rows)==count
    assert '22 맵 순서와 운동 배치' in pdf.pages[21].extract_text()
    assert '23 플레이 진행에 따른 속도와 안내' in pdf.pages[22].extract_text()
    assert '24 조작 기준과 플레이테스트' in pdf.pages[23].extract_text()
    body='\n'.join(p.text for p in document.paragraphs)
    body+='\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert not re.search(r'\d+(?:\.\d+)?\s*m(?:\b|/s)|\d+(?:\.\d+)?\s*미터',body)
    for phrase in ['맵 길이와 장애물 사이 간격은 아직 미정','거리와 간격은 실제 플레이를 보며 정한다','점프가 우선','V10 촬영본']:
        assert phrase in body,phrase
gallery = root / 'docs/previews/ArtV10/index.html'
refs = re.findall(r'<img[^>]+src=["\']([^"\']+)', gallery.read_text(encoding='utf-8'))
assert refs
for ref in refs:
    assert (gallery.parent / ref).exists(), ref
print(f'SCREEN_PLAN_VERIFICATION_PASS: {args.version} / {expected} rendered pages / 20 photos / {10 if args.version in ("V12","V13") else 0} diagrams / {len(document.tables)} tables / {len(refs)} gallery images')

"""Revise the retained V10 screen plan into the current V11 gameplay plan."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/무한옥_화면별_게임기획서_V10.docx'
OUTPUT = ROOT / 'docs/무한옥_화면별_게임기획서_V11.docx'
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)

revisions = {
    1: '화면 흐름과 맵별 운동 점진 가속 설계',
    2: '2026년 10월 5일   |   기획 V11   |   화면 01 타이틀',
    5: '무한옥은 광산, 수레 철로, 계단, 감옥, 사육장을 달리는 운동 게임이다. V11은 도보 맵의 점프와 슬라이드, 이동 거리에 따른 가속을 반영한다.',
    11: '게임은 16대9 비율을 유지한다. 사진은 V10 촬영본으로 공간과 연출의 참고 자료다. 현행 장애물 배치와 속도 규칙은 본문과 22~24장의 표를 따른다.',
    38: '전환 중에는 이동과 운동 시간을 멈춘다. 완료 뒤 좌우 이동, 점프, 슬라이드, 하이니 입력을 받으며 기본 9 m/s로 달리기를 시작한다.',
    42: '전환이 끝나면 좌우 이동, 점프, 슬라이드, 하이니 입력을 받는다. 시작과 재시작은 기본 9 m/s이며, 이후에는 실제 이동 거리로 속도를 계산한다.',
    45: '05 광산 달리기와 운동 안내',
    49: '광산에서 1인칭 달리기를 시작한다. 자동으로 전진하며 점프와 슬라이드로 장애물을 통과한다. 이동할수록 속도가 올라 같은 구간의 대응 시간이 짧아진다.',
    50: '맵별 동작',
    51: '일반 도보 맵은 로컬 7 m의 낮은 상자를 점프하고 19 m의 파이프 아래로 슬라이드한다. 세 레인 모두 같은 동작을 요구한다. 첫 광산 진입에서는 점프 행을 생략해 준비 시간을 준다.',
    54: 'HUD와 다음 구간',
    55: 'HUD에는 거리, 코인, 현재 속도와 다음 동작을 표시한다. 안내 거리는 기본 속도의 약 1.35초 진행량에 맞춰 늘어난다. 몸무게 입력 시 예상 kcal도 표시한다. 앞쪽 승강장에서는 수레로 자동 탑승한다.',
    60: '수레에 탑승하면 단선 절벽 철로를 따라 자동으로 전진한다. 카메라는 뒤쪽 3인칭으로 바뀌며, 도보와 같은 거리별 가속을 적용한다.',
    64: '수레에서는 좌우 기울이기와 숙이기를 사용한다. 점프와 하이니 입력은 받지 않는다. 일반 도보 장애물을 배치하지 않고 붕괴 철로에서 반대쪽 몸 기울이기를 요구한다.',
    66: '철로를 따라 수레 본체와 원숭이의 몸이 함께 기울어진다. 하차하면 1인칭과 세 레인 조작으로 돌아오며 누적 거리와 가속을 유지한다.',
    86: '현재 화면의 카메라 위치를 다음 시점의 시작점으로 사용한다. 하차 후에는 1인칭 손이 다시 나타나고, 기본 속도를 초기화하지 않은 채 도보 조작을 이어 간다.',
    88: '하차 후 시설 연결 구간에서 점프와 슬라이드를 사용한다. 이어지는 계단은 점프, 하이니, 문 아래 슬라이드 순서로 통과한다.',
    97: '계단의 로컬 3 m에서 상자를 점프하고 12 m에서 하이니 장애물을 통과한다. 점프만으로 하이니를 대신할 수 없다. 낮은 발은 경로 높이에 맞추며 공중에서는 접지 보정을 해제한다.',
    99: '계단 끝 로컬 19 m에서 감옥 문 아래로 슬라이드한다. 계단 높이는 3 m 상승하며, 다음 감옥 복도와 바닥 및 진행 방향을 이어 놓는다.',
    117: '숙이기 애니메이션과 낮아지는 카메라로 슬라이드를 표현한다. CROUCH의 기본 유지 시간은 마지막 입력부터 0.8초다. S 또는 아래 방향키를 누르는 동안 재입력으로 연장되며 몸동작도 같은 규칙을 사용한다.',
    121: '문 통과 순간에 슬라이드 중이면 성공한다. 바닥에서 점프하면 진행 중인 슬라이드를 해제하고 점프를 우선한다. 공중 연속 점프와 공중 슬라이드는 허용하지 않는다.',
    132: '감옥 복도에서는 로컬 7 m의 상자를 점프하고 19 m의 파이프 아래로 슬라이드한다. 23 m의 닫히는 문은 한 레인만 막으므로 다른 두 레인으로 이동해 피한다.',
    154: '점프와 수레에서는 바닥 접지 보정을 해제한다. 진행과 감속에 맞춰 달리기 재생 속도를 조절한다. 캐릭터의 실제 위치 이동은 게임 규칙이 담당하고 뼈대 동작은 애니메이션이 담당한다.',
    163: '보안 출구도 7 m 점프와 19 m 슬라이드 행을 갖는다. 23 m의 닫히는 문은 옆 레인으로 피한다. 운동 행에서 차선만 바꿔 해당 동작을 생략할 수 없도록 한다.',
    174: '사육장과 외부 내리막에도 7 m 점프와 19 m 슬라이드 행을 배치한다. 바나나와 물웅덩이는 옆 차선으로 피하거나 점프로 넘는다. 고체 충돌은 종료, 바닥 아이템 접촉은 감속으로 구분한다.',
    176: '사육장을 지나면 광산부터 같은 경로를 반복한다. 매 반복에서 실제 이동 거리를 이어 누적하므로 기본 속도가 계속 올라가며, 960 m부터 17 m/s를 유지한다.',
    183: '바나나 첫 접촉은 2.4초 동안 당시 기본 속도의 50%로 감속한다. 접촉 횟수는 이번 플레이에 남는다. 감속 중에도 실제로 이동한 거리에 따라 기본 속도는 계속 증가한다.',
    187: '감속 중 카메라에 짧은 좌우 흔들림을 적용하고 달리기 재생 속도를 낮춘다. 감속이 끝나면 현재 누적 거리의 기본 속도로 회복한다. 예를 들어 480 m 부근에서는 약 6.5 m/s로 느려진다.',
    198: '잡힘 연출 후 결과 화면으로 전환한다. 다시 시작하면 바나나 횟수, 감속, 추격, 점수와 이동 거리를 초기화한다. 새 플레이는 다시 기본 9 m/s에서 시작한다.',
    207: '자동 전진, 가속을 결정하는 이동 거리, 추격과 감속 시간, 운동 시간을 함께 멈춘다. HUD의 현재 속도는 0 m/s다. 재개 시 누적 거리를 유지하며 대기 중 입력을 비운다.',
    209: '재개 버튼으로 같은 위치에서 이어 한다. 정지 시간에는 거리와 예상 칼로리가 증가하지 않는다. 슬라이드나 점프 타이머도 정지 전 상태를 유지한다.',
    220: '재시작하면 타이틀로 돌아가 몸무게 선택과 로딩을 거친다. 거리와 속도, 바나나 횟수, 장애물 통과 상태와 문을 초기화한다. 씬을 다시 불러오지 않고 맵 객체를 재사용한다.',
    222: '낮은 상자는 점프, 머리 위 파이프는 슬라이드, 계단형 운동 장애물은 하이니로 통과한다. 세 레인 전체의 운동 행은 같은 동작을 요구하며, 차선을 막는 문은 다른 레인을 남긴다.',
    224: 'A와 D 또는 좌우 방향키로 이동한다. Space로 점프, S 또는 아래 방향키로 슬라이드, W 또는 위 방향키로 하이니를 사용한다. 슬라이드 중 바닥 점프는 점프를 우선하고 수레에서는 점프를 받지 않는다.',
    226: '점프와 슬라이드가 각 도보 맵에서 연결되는지 확인한다. 계단은 점프, 하이니, 닫히는 문 슬라이드 순서다. 수레의 기울이기와 독방 원숭이의 배치 및 접지는 기존 공간 기준을 유지한다.',
}
for index, text in revisions.items():
    assert doc.paragraphs[index].text, (index, 'Empty source slot')
    assert doc.paragraphs[index].style.name not in ('Heading 1', 'Heading 2') or len(text) < 35, (index, 'Heading capacity')
    doc.paragraphs[index].text = text
for p in doc.paragraphs:
    if p.style.name == 'Caption' and p.text.startswith('화면 '):
        p.text = p.text.replace('실제 Unity 실행 화면', 'V10 실제 화면 참고')
    if p.style.name == 'Caption' and p.text.startswith('그림 '):
        p.text = p.text.replace('실제 플레이 화면', 'V10 실제 플레이 화면')
        p.text = p.text.replace('제작 검수용 근접 카메라', 'V10 제작 검수용 근접 카메라')
for footer in doc.sections[0].footer.paragraphs:
    for run in footer.runs:
        run.text = run.text.replace('V10', 'V11')

doc.tables[0].cell(1, 0).text = '낮은 상자'
doc.tables[0].cell(2, 1).text = '몸을 낮춰 슬라이드'
doc.tables[0].cell(4, 1).text = '슬라이드로 아래 통과'

def heading(number, title, intro):
    p = doc.add_paragraph(f'{number:02d} {title}', 'Heading 1')
    p.paragraph_format.page_break_before = True
    doc.add_paragraph(intro)

def block(title, text):
    doc.add_paragraph(title, 'Heading 2')
    doc.add_paragraph(text)

def table(headers, records, widths, centered=()):
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for cell, text, width in zip(result.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = text
    for values in records:
        for cell, text, width in zip(result.add_row().cells, values, widths):
            cell.width = Inches(width)
            cell.text = text
    header = OxmlElement('w:tblHeader')
    result.rows[0]._tr.get_or_add_trPr().append(header)
    for row_index, row in enumerate(result.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                value = OxmlElement('w:' + side)
                value.set(qn('w:w'), '95' if side in ('top', 'bottom') else '110')
                value.set(qn('w:type'), 'dxa')
                margins.append(value)
            cell._tc.get_or_add_tcPr().append(margins)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7ECF1')
                cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.12
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in centered else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    result._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return result

heading(22, '맵 순서와 운동 배치', '본편은 아래 10개 청크를 순서대로 반복한다. 청크 길이는 24 m이며 위치는 청크 입구 기준이다. 도보 운동 행은 세 레인 모두 같은 동작을 요구한다. 전체 13종 프리팹 중 도보 11종에는 점프와 슬라이드를, 수레 2종에는 철로 위험을 준비했다.')
table(['순서', '맵 구간', '점프', '슬라이드', '추가 행동'], [
    ('1', '광산 시작', '7 m', '19 m', '첫 진입만 점프 생략'),
    ('2', '수레 철로 전반', '없음', '도보 행 없음', '8 m와 18 m에서 기울이기'),
    ('3', '수레 철로 후반', '없음', '도보 행 없음', '8 m와 18 m에서 기울이기'),
    ('4', '시설 연결', '7 m', '19 m', '하차 후 1인칭으로 복귀'),
    ('5', '시설 계단', '3 m', '19 m', '12 m 하이니와 입구 문'),
    ('6', '감옥 복도 A', '7 m', '19 m', '23 m 문을 옆 레인으로 회피'),
    ('7', '감옥 복도 B', '7 m', '19 m', '23 m 문을 옆 레인으로 회피'),
    ('8', '보안 출구', '7 m', '19 m', '23 m 문을 옆 레인으로 회피'),
    ('9', '외부 내리막', '7 m', '19 m', '높이 3 m에서 0 m로 연결'),
    ('10', '동물 사육장', '7 m', '19 m', '바닥 아이템 회피 후 광산 반복'),
], [.45, 1.35, .62, 1.13, 3.55], centered=(0, 2, 3))
block('수레 조작의 예외', '수레는 좌우 기울이기와 숙이기를 허용하지만 점프와 하이니를 받지 않는다. 일반 상자와 파이프 운동 행은 배치하지 않는다. 붕괴한 쪽의 반대 방향으로 몸을 기울여 남은 레일의 균형을 잡는다.')
block('계단과 감옥 문', '계단 끝의 입구 문은 세 레인을 막아 반드시 슬라이드해야 한다. 통과 뒤 3 m에서 완전히 닫힌다. 감옥 복도의 별도 문은 한 레인만 막아 다른 두 레인으로 피할 수 있다.')

heading(23, '진행 거리에 따른 속도와 안내', '속도는 실제로 이동한 누적 거리로 결정한다. 도보와 수레 모두 기본 9 m/s에서 출발하며 거리 120 m마다 1 m/s씩 연속으로 증가한다. 최대 속도는 17 m/s다.')
block('가속 계산', '기본 속도는 9에 누적 이동 거리 나누기 120을 더한 값이다. 이동 거리에 비례해 매 프레임 조금씩 증가하며 17 m/s에 도달하면 더 증가하지 않는다.')
table(['누적 이동 거리', '기본 속도', '50퍼센트 감속 시'], [
    ('0 m', '9.0 m/s', '4.5 m/s'),
    ('120 m', '10.0 m/s', '5.0 m/s'),
    ('240 m', '11.0 m/s', '5.5 m/s'),
    ('480 m', '13.0 m/s', '6.5 m/s'),
    ('960 m 이상', '17.0 m/s', '8.5 m/s'),
], [2.3, 2.2, 2.6], centered=(0, 1, 2))
block('감속과 속도 회복', '바나나와 물웅덩이는 2.4초 동안 기본 속도의 50%로 감속한다. 이때 실제 이동한 거리도 가속 계산에 포함된다. 감속이 끝나면 감속 전 수치가 아니라 현재 누적 거리의 기본 속도로 돌아온다.')
block('HUD와 다음 동작 안내', 'HUD의 속도는 현재 감속까지 반영한 m/s 값이다. 다음 동작은 점프, 슬라이드, 하이니, 문 회피 또는 수레 기울이기로 안내한다. 도보 안내 거리는 기본 속도의 약 1.35초 진행량에 맞춰 커지며, 레인 차단 문은 플레이어가 향하는 레인에 있을 때 안내한다.')
block('정지와 재시작', '일시정지와 잡힘에서는 이동을 멈추고 현재 속도를 0 m/s로 표시한다. 재개는 기존 누적 거리의 속도로 이어 간다. 재시작은 이동 거리를 0 m로 초기화해 다시 9 m/s로 출발한다. 반복 맵 진입이나 수레 하차로는 거리와 가속을 초기화하지 않는다.')

heading(24, '조작 기준과 구현 검수', '키보드와 웹캠 몸동작은 공통 게임 명령을 사용한다. 구현된 판정과 자동 검수 결과를 기준으로 현재 기능을 정의하며, 실제 몸동작의 난이도는 별도로 확인한다.')
table(['동작', '키보드', '판정과 기본 설정'], [
    ('레인 이동', 'A D 또는 좌우 키', '레인 간격 2.2 m, 이동 시간 0.20초'),
    ('점프', 'Space', '바닥에서 시작, 높이 1.8 m, 공중 연속 점프 금지'),
    ('슬라이드', 'S 또는 아래 키', '기본 0.8초, 키를 누르는 동안 입력 시간 갱신'),
    ('하이니', 'W 또는 위 키', '기본 유지 0.9초, 점프로 계단 판정 대체 불가'),
    ('수레 기울이기', 'A D 또는 좌우 키', '붕괴 쪽 반대로 기울이기, 점프와 하이니 금지'),
], [1.15, 1.7, 4.25])
block('연속 동작과 실패', '바닥에서 점프하면 진행 중인 슬라이드를 해제한다. S를 누른 채 Space를 눌러도 점프가 우선한다. 공중 슬라이드는 받지 않는다. 고체 장애물은 한 번 충돌하면 종료하며, 바나나는 같은 플레이에서 두 번째 접촉 후 약 1.2초 추격으로 종료한다. 물웅덩이는 바나나 횟수에 포함하지 않는다.')
block('현재 검수 결과', 'Unity 6000.3.10f1에서 EditMode 49개와 PlayMode 23개를 모두 통과했다. 최고 속도 17 m/s에서 실제 풀링된 맵 두 바퀴를 30fps로 통과했으며 점프, 슬라이드, 하이니, 문 회피와 수레 기울이기를 확인했다. C# 컴파일 오류는 없으며 Windows 빌드와 실행 후 재시작 검증도 통과했다.')
block('자료와 후속 확인', '사진 20장은 V10 실행 화면이며 새 속도 표시와 운동 행을 재촬영한 사진은 아니다. 현행 규칙은 실제 코드와 MAP_EXERCISES_SPEED.md를 기준으로 작성했다. 자동 최고 속도 조작 검수는 바닥 아이템과 실제 웹캠 인식 정확도를 포함하지 않는다. 다음 검수는 실제 운동으로 최대 속도와 안내 시간을 조정하고 최신 HUD 및 운동 장면을 재촬영하는 것이다.')

doc.core_properties.title = '무한옥 화면별 게임 기획서 V11'
doc.core_properties.subject = '맵별 점프 슬라이드 점진 가속 조작과 화면 흐름'
doc.core_properties.comments = ''
settings = doc.settings.element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')

buffer = BytesIO()
doc.save(buffer)
# Preserve all unrelated package parts, including images, styles and relationships.
editable = {'word/document.xml', 'word/footer1.xml', 'word/settings.xml', 'docProps/core.xml'}
with zipfile.ZipFile(SOURCE) as original, zipfile.ZipFile(buffer) as revised:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(OUTPUT, 'w', zipfile.ZIP_DEFLATED) as destination:
        for item in original.infolist():
            destination.writestr(deepcopy(item), revised.read(item.filename) if item.filename in editable else original.read(item.filename))
    with zipfile.ZipFile(OUTPUT) as final:
        for name in original.namelist():
            if name not in editable:
                assert original.read(name) == final.read(name), name
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
result = Document(OUTPUT)
assert len(result.inline_shapes) == 20 and len(result.tables) == 4
audit = {'source_sha256': source_hash, 'source_preserved': True, 'images': 20, 'tables': 4,
         'intended_pages': 24, 'edited_package_parts': sorted(editable), 'output': str(OUTPUT)}
qa = ROOT / 'docs/document-qa/V11'
qa.mkdir(parents=True, exist_ok=True)
(qa / 'document-audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'SCREEN_PLAN_V11_CREATED: {OUTPUT} / 20 retained photos / 4 tables / 24 intended pages')

"""Remove every footer and illustrate obstacle avoidance in the V12 screen plan."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import math
import zipfile
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/무한옥_화면별_게임기획서_V11.docx'
OUTPUT = ROOT / 'docs/무한옥_화면별_게임기획서_V12.docx'
ART = ROOT / 'docs/document-qa/V12/obstacle-diagrams'
ART.mkdir(parents=True, exist_ok=True)
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
FONT = 'C:/Windows/Fonts/malgun.ttf'
INK, SAFE, RED = '#263441', '#087C63', '#C94037'

def draw_diagram(kind, key, action):
    im = Image.new('RGB', (1400, 370), 'white')
    d = ImageDraw.Draw(im)
    def label(x, y, text, size=27, color=INK):
        d.text((x, y), text, font=ImageFont.truetype(FONT, size), fill=color)
    def line(points, color=INK, width=8):
        d.line(points, fill=color, width=width, joint='curve')
    def arrow(points, color=SAFE):
        line(points, color, 8)
        x,y=points[-1]; px,py=points[-2]
        a=math.atan2(y-py,x-px)
        d.polygon([(x,y),(x-22*math.cos(a-.5),y-22*math.sin(a-.5)),
                   (x-22*math.cos(a+.5),y-22*math.sin(a+.5))], fill=color)
    def person(x, y, pose='stand'):
        # y is the foot height; all figures represent the player, not a captured model.
        if pose == 'crouch':
            d.ellipse((x-14,y-103,x+18,y-71),outline=SAFE,width=7)
            line([(x,y-68),(x+54,y-47),(x+13,y-24),(x+53,y-4)],SAFE)
            line([(x+49,y-44),(x+72,y-12)],SAFE)
            line([(x+17,y-63),(x+48,y-69)],SAFE)
        elif pose == 'lean':
            d.ellipse((x+21,y-152,x+55,y-118),outline=SAFE,width=7)
            line([(x+36,y-114),(x,y-65),(x-15,y)],SAFE)
            line([(x,y-65),(x+43,y)],SAFE)
            line([(x+27,y-99),(x-28,y-98)],SAFE)
        else:
            d.ellipse((x-17,y-153,x+17,y-119),outline=SAFE,width=7)
            line([(x,y-115),(x,y-65)],SAFE)
            line([(x,y-98),(x-36,y-76)],SAFE)
            line([(x,y-98),(x+35,y-83)],SAFE)
            if pose == 'knee':
                line([(x,y-65),(x+48,y-89),(x+55,y-43)],SAFE)
                line([(x,y-65),(x-12,y)],SAFE)
                arrow([(x+79,y-14),(x+79,y-80)])
            elif pose == 'jump':
                line([(x,y-65),(x-38,y-49),(x-27,y-12)],SAFE)
                line([(x,y-65),(x+35,y-45),(x+18,y-13)],SAFE)
            else:
                line([(x,y-65),(x-27,y)],SAFE)
                line([(x,y-65),(x+27,y)],SAFE)
    label(35, 8, '장애물 모습', 27)
    label(790, 8, '피하는 동작', 27)
    line([(700,55),(700,315)], '#DAE0E5', 2)
    line([(40,282),(650,282)], '#788591', 3)
    line([(775,282),(1350,282)], '#788591', 3)
    if kind == 'crate':
        d.rectangle((250,169,450,280),fill='#B27A46',outline='#684829',width=6)
        for x in [275,320,365,410]: line([(x,174),(x,276)], '#825832', 3)
        line([(260,180),(440,270)], '#D4A777', 11)
        line([(440,180),(260,270)], '#D4A777', 11)
        d.rectangle((1040,217,1190,280),fill='#B27A46',outline='#684829',width=4)
        person(1010,187,'jump')
        arrow([(825,230),(885,126),(965,77),(1060,83),(1225,222)])
        label(120,303,'세 레인 모두 낮은 상자',24)
    elif kind == 'pipe':
        for x in [220,490]:
            d.rectangle((x,95,x+24,280),fill='#657785')
        d.rounded_rectangle((207,82,530,120),radius=18,fill='#536673')
        for x in [275,430]: d.rectangle((x,76,x+29,126),fill='#A5B5BD')
        d.rounded_rectangle((840,118,1295,149),radius=12,fill='#536673')
        for x in [848,1267]: d.rectangle((x,140,x+19,281),fill='#657785')
        person(1040,278,'crouch')
        arrow([(916,222),(1185,222)])
        label(130,303,'머리 높이의 가로 파이프',24)
    elif kind == 'stairs':
        for i in range(4):
            x=215+i*64; top=260-i*39
            d.rectangle((x,top,x+65,280),fill='#9BABB4',outline='#62717B',width=3)
            line([(x,top),(x+65,top)], '#C49C4D',5)
        person(1040,281,'knee')
        label(160,303,'작은 계단 모양 운동 행',24)
    elif kind == 'slidegate':
        for x in [200,510]: d.rectangle((x,62,x+22,282),fill='#495963')
        d.rectangle((222,69,509,181),fill='#8999A6',outline='#44545F',width=4)
        for y in range(80,178,22): line([(230,y),(500,y)], '#556976',4)
        d.rectangle((240,114,490,144),fill='#E2CFAC')
        for x in range(245,470,48): line([(x,144),(x+21,114)],RED,14)
        label(225,215,'아래 틈',24,SAFE)
        d.rectangle((850,64,1290,155),fill='#8999A6',outline='#44545F',width=4)
        person(1040,279,'crouch'); arrow([(925,227),(1190,227)])
        label(160,303,'계단 끝의 전폭 입구 문',24)
    elif kind in ('door','barrier','laser'):
        for x in [170,325,480]:
            d.rectangle((x,70,x+135,280),outline='#C4CFD6',width=3)
        if kind=='door':
            d.rectangle((332,78,452,280),fill='#95A4AD',outline='#44545F',width=4)
            for y in range(92,270,28): line([(337,y),(448,y)],'#566B79',4)
            d.rectangle((337,126,448,162),fill='#E5D2B5')
            for x in range(346,435,28): line([(x,158),(x+13,131)],RED,9)
        elif kind=='barrier':
            d.rectangle((333,143,451,271),fill='#E5D2B5',outline='#44545F',width=4)
            for x in range(337,422,28): line([(x,246),(x+25,154)],RED,11)
        else:
            for x in [330,446]:d.rectangle((x,84,x+12,280),fill='#44545F')
            for y in [135,190,245]:line([(341,y),(446,y)],RED,8)
        label(190,301,'안전',24,SAFE); label(336,301,'위험',24,RED); label(493,301,'안전',24,SAFE)
        for x in [815,990,1165]: d.rectangle((x,73,x+153,280),outline='#C4CFD6',width=3)
        person(887,280)
        d.rectangle((1011,85,1132,278),fill='#DFDFDC')
        arrow([(1070,211),(924,211)])
    elif kind == 'rail':
        for x in [242,440]:line([(x,79),(x,278)],'#657785',17)
        for y in range(100,279,38):line([(213,y),(468,y)],'#8C653E',15)
        d.rectangle((229,171,256,251),fill='white')
        line([(225,170),(249,188),(230,206)],RED,5)
        label(93,194,'왼쪽 붕괴',24,RED)
        d.rectangle((971,218,1170,279),fill='#82949E',outline='#44545F',width=4)
        person(1067,258,'lean')
        arrow([(1080,75),(1205,75)])
        label(186,303,'한쪽 레일이 끊긴 구간',24)
    elif kind == 'banana':
        d.polygon([(350,189),(291,218),(215,241),(230,269),(295,255),(350,219),
                   (415,260),(482,271),(492,244),(416,222),(367,190)],fill='#F4C443',outline='#9B7321')
        d.polygon([(350,198),(355,258),(378,278),(398,261),(380,212)],fill='#EFB92E')
        line([(350,189),(358,175)],'#685331',10)
        person(951,280);arrow([(1090,210),(994,210)])
        d.polygon([(1150,242),(1122,270),(1190,273),(1170,242)],fill='#F4C443')
        label(162,303,'바닥의 노란 바나나 껍질',24)
    elif kind == 'puddle':
        d.ellipse((200,187,512,279),fill='#53899F',outline='#2A5C78',width=5)
        d.arc((240,203,470,260),190,320,fill='#A8D1E3',width=4)
        d.ellipse((1080,230,1250,280),fill='#53899F',outline='#2A5C78',width=3)
        person(936,280);arrow([(1090,198),(975,198)])
        label(179,303,'바닥의 푸른 물웅덩이',24)
    label(800,320,f'{key}   {action}',27,SAFE)
    path=ART/f'{kind}.png'; im.save(path)
    return path

doc=Document(SOURCE)
# Clear default, first-page and even-page footer content, then detach every footer.
for section in doc.sections:
    for footer in (section.footer,section.first_page_footer,section.even_page_footer):
        for child in list(footer._element): footer._element.remove(child)
        footer._element.append(OxmlElement('w:p'))
    for ref in list(section._sectPr.findall(qn('w:footerReference'))):
        section._sectPr.remove(ref)

for p in doc.paragraphs:
    if '기획 V11' in p.text: p.text=p.text.replace('기획 V11','기획 V12')
    if 'V11은' in p.text: p.text=p.text.replace('V11은','V12는')
    if p.text=='화면 흐름과 맵별 운동 점진 가속 설계':
        p.text='화면 흐름과 장애물 회피 점진 가속 설계'
    if '22~24장의 표' in p.text:
        p.text=p.text.replace('22~24장의 표','22~24장의 표와 25~29장의 회피 그림')
    if p.text.startswith('낮은 상자는 점프, 머리 위 파이프는 슬라이드'):
        p.text='상자는 점프, 파이프와 입구 문은 슬라이드, 계단 운동 행은 하이니로 통과한다. 레인 차단 문은 옆으로 피한다. 장애물 모습과 몸동작은 25~29장의 그림으로 확인한다.'
    if p.text.startswith('사진 20장은 V10 실행 화면이며'):
        p.text='사진 20장은 V10 참고 화면이며 회피 설명도 10장은 장애물과 조작을 보여 준다. 현행 규칙은 실제 코드와 MAP_EXERCISES_SPEED.md를 기준으로 작성했다. 자동 최고 속도 검수는 바닥 아이템과 실제 웹캠 인식 정확도를 포함하지 않는다. 다음 검수는 실제 운동으로 속도와 안내 시간을 조정하고 최신 HUD 및 운동 장면을 재촬영하는 것이다.'

route=doc.tables[1]
route.cell(0,2).text='점프\nSpace'
route.cell(0,3).text='슬라이드\nS / ↓'
for i in [1,4,6,7,8,9,10]:
    route.cell(i,2).text='7 m'
    route.cell(i,3).text='19 m'
route.cell(5,2).text='3 m'
route.cell(5,3).text='19 m 문'
route.cell(5,4).text='12 m 계단은 W / ↑ 하이니'
for i in [2,3]:route.cell(i,4).text='8 m·18 m 붕괴 반대쪽 기울이기'
for i in [6,7,8]:route.cell(i,4).text='23 m 차단 문은 A / D로 옆 레인 회피'
route.cell(10,4).text='바나나·물웅덩이는 옆 레인 또는 점프'
for row_index,row in enumerate(route.rows):
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.line_spacing=1.12
            for run in p.runs:
                run.font.size=Pt(10.5)
                run.font.bold=row_index==0
for p in doc.paragraphs:
    if p.text.startswith('수레는 좌우 기울이기와 숙이기를 허용하지만'):
        p.text='수레는 붕괴 반대쪽으로 기울여 통과한다. 점프와 하이니는 받지 않으며 일반 상자·파이프 운동 행은 배치하지 않는다. 숙이기는 허용한다.'
    if p.text.startswith('계단 끝의 입구 문은 세 레인을 막아'):
        p.text='계단 끝 입구 문은 세 레인을 막아 슬라이드로 아래 틈을 통과한다. 감옥 복도 문은 바닥까지 닫혀 옆 레인으로 피한다. 두 문의 회피 방법은 다르다.'

records=[
 ('crate','낮은 상자','Space','점프하여 넘기',
  '일반 도보 맵의 7 m, 계단의 3 m에서 나온다. 낮은 나무 상자가 세 레인에 있으므로 차선 변경으로 생략할 수 없다. 첫 광산 진입에서는 상자 행을 생략한다.',
  '상자에 닿기 전에 몸을 점프하거나 Space를 눌러 통과 순간 공중에 있어야 한다. 슬라이드로 통과할 수 없고, 상자와 충돌하면 즉시 종료한다.'),
 ('pipe','머리 위 파이프','S / ↓','몸을 낮춰 슬라이드',
  '일반 도보 맵의 19 m에서 나온다. 회색 가로 파이프와 양쪽 지지대가 세 레인에 같은 운동 행을 만든다. 계단의 19 m에는 파이프 대신 입구 문이 나온다.',
  '파이프에 도달하기 전에 몸을 숙이거나 S / ↓를 눌러 통과 순간 슬라이드 상태를 유지한다. 키를 누르고 있으면 유지 시간이 갱신된다. 서 있거나 점프해 부딪히면 종료한다.'),
 ('stairs','계단형 하이니 장애물','W / ↑','무릎 높이 들기',
  '시설 계단의 12 m에서 나온다. 작은 계단 모양의 운동 행이 세 레인에 있으므로 다른 레인으로 이동해도 하이니가 필요하다.',
  '무릎을 높이 드는 하이니 동작을 하거나 W / ↑를 눌러 통과 순간 하이니 상태여야 한다. 기본 유지 시간은 0.9초다. 일반 점프만으로 대신할 수 없으며 판정에 실패하면 종료한다.'),
 ('slidegate','계단 끝 입구 문','S / ↓','아래 틈으로 통과',
  '시설 계단의 19 m에 있다. 입구 문이 세 레인 전체를 막지만 처음에는 바닥 위에 낮은 틈을 남긴다. 붉은 경고 줄무늬와 문 아래 빈 공간을 확인한다.',
  '도달하기 전에 몸을 낮추거나 S / ↓를 눌러 문 아래로 슬라이드한다. 옆 레인이나 점프로 대신할 수 없다. 통과 뒤 3 m에서 문이 완전히 닫힌다. 서서 닿으면 종료한다.'),
 ('door','감옥 복도의 차단 문','A / D','빈 레인으로 이동',
  '감옥 복도 A·B와 보안 출구의 23 m에서 나온다. 붉은 경고등이 켜진 철제 문이 한 레인을 바닥까지 막고 다른 두 레인은 비워 둔다.',
  '문이 막는 레인을 확인한 뒤 몸을 좌우로 움직이거나 A / D 또는 좌우 키로 빈 레인에 들어간다. 아래 틈이 없어 슬라이드로 지나갈 수 없다. 문에 닿으면 즉시 종료한다.'),
 ('rail','수레 철로 붕괴','A / D','붕괴 반대로 기울이기',
  '두 수레 맵의 8 m와 18 m에서 한쪽 레일이 끊긴다. 그림은 왼쪽 레일이 무너진 경우다. 오른쪽 붕괴에서는 방향을 반대로 적용한다.',
  '왼쪽이 무너지면 오른쪽으로 기울이거나 D / →를 누른다. 오른쪽이 무너지면 왼쪽으로 기울이거나 A / ←를 누른다. 수레에서는 차선을 바꾸거나 점프하지 않는다. 잘못 기울이면 추락해 종료한다.'),
 ('banana','바나나 껍질','A / D 또는 Space','옆으로 피하거나 점프',
  '동물 사육장 바닥에서 노란 껍질 형태로 나온다. 코인과 달리 바닥에 낮게 놓인다. 해당 레인을 미리 피하거나 점프로 넘는다.',
  '몸을 좌우로 이동하거나 A / D로 빈 레인에 들어간다. 같은 레인에서는 몸 점프 또는 Space로 넘는다. 첫 접촉은 2.4초간 50% 감속하며, 같은 플레이의 두 번째 접촉은 약 1.2초 추격 후 종료한다.'),
 ('puddle','물웅덩이','A / D 또는 Space','옆으로 피하거나 점프',
  '동물 사육장 바닥에 푸른 반사면으로 나온다. 바닥에 넓고 낮게 퍼져 있어 정면에서 물기가 보이는 레인을 확인한다.',
  '몸을 좌우로 이동하거나 A / D로 다른 레인에 들어간다. 몸 점프 또는 Space로 넘을 수도 있다. 닿으면 2.4초간 50% 감속한다. 물웅덩이 접촉은 바나나 횟수에 더하지 않는다.'),
 ('barrier','차선 차단벽','A / D','빈 레인으로 이동',
  '붉은 줄무늬가 있는 차단판으로 한 레인을 막는 보조 장애물이다. 현재 10개 본편 경로의 운동 행에는 사용하지 않으며 다른 맵 변형에 배치할 때 적용한다.',
  '차단판이 없는 옆 레인으로 몸을 움직이거나 A / D를 누른다. 낮은 상자와 달리 점프나 슬라이드로 통과하는 대상으로 처리하지 않는다. 충돌하면 종료한다.'),
 ('laser','레이저 차단기','A / D','빈 레인으로 이동',
  '양쪽 기둥 사이의 붉은 가로 빛이 레인을 막는 보조 장애물이다. 현재 10개 본편 경로에는 사용하지 않으며 다른 맵 변형에 배치할 때 적용한다.',
  '붉은 빛이 없는 옆 레인으로 몸을 움직이거나 A / D를 누른다. 점프나 숙이기로 빛 사이를 지나가는 대상으로 처리하지 않는다. 같은 레인으로 닿으면 종료한다.'),
]
pages=[('점프와 파이프 슬라이드','도보 운동 행은 세 레인 모두 같은 동작을 요구한다. 그림에서 장애물 형태를 확인하고 통과 순간에 맞춰 동작한다.'),
       ('하이니와 입구 문 슬라이드','계단에서는 상자 점프 다음에 하이니, 입구 문 아래 슬라이드를 이어 간다. 하이니와 점프는 서로 다른 판정이다.'),
       ('문 회피와 수레 균형','감옥 차단 문은 빈 레인으로 이동하고 수레 철로는 남은 레일 쪽으로 기울인다. 수레 조작은 도보의 레인 이동과 다르다.'),
       ('바닥 아이템 피하기','바나나와 물웅덩이는 옆 레인으로 피하거나 점프로 넘는다. 접촉 시 감속은 같지만 바나나만 누적 추격으로 연결된다.'),
       ('보조 장애물의 회피 기준','차단벽과 레이저는 기존 판정 표의 보조 유형이다. 현재 본편 경로의 운동 행과 구분하며 배치하는 경우 아래 회피 기준을 적용한다.')]
for page_idx,(title,intro) in enumerate(pages):
    p=doc.add_paragraph(f'{25+page_idx:02d} {title}','Heading 1')
    p.paragraph_format.page_break_before=True
    doc.add_paragraph(intro)
    for index in range(page_idx*2,page_idx*2+2):
        kind,name,key,action,where,how=records[index]
        doc.add_paragraph(name,'Heading 2')
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        run=p.add_run();shape=run.add_picture(str(draw_diagram(kind,key,action)),width=Inches(7.0))
        shape._inline.docPr.set('descr',f'{name}의 모습과 {action} 동작을 보여 주는 회피 설명도')
        p=doc.add_paragraph(f'그림 {21+index} {name} 회피 설명도','Caption')
        p.paragraph_format.space_after=Pt(4)
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
        p.add_run('등장과 식별  ').bold=True;p.add_run(where)
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7)
        p.add_run('회피와 실패  ').bold=True;p.add_run(how)

doc.core_properties.title='무한옥 화면별 게임 기획서 V12'
doc.core_properties.subject='장애물 모습과 회피 동작 맵 운동 점진 가속'
buffer=BytesIO();doc.save(buffer)
# Keep the retained template, original media and unrelated package parts byte-for-byte.
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
with zipfile.ZipFile(SOURCE) as original,zipfile.ZipFile(buffer) as revised:
    original_names=set(original.namelist())
    editable.update(name for name in original_names if name.startswith('word/footer') and name.endswith('.xml'))
    with zipfile.ZipFile(OUTPUT,'w',zipfile.ZIP_DEFLATED) as destination:
        for item in original.infolist():
            destination.writestr(deepcopy(item),revised.read(item.filename) if item.filename in editable else original.read(item.filename))
        for name in revised.namelist():
            if name not in original_names:destination.writestr(name,revised.read(name))
    with zipfile.ZipFile(OUTPUT) as final:
        for name in original_names-editable:assert original.read(name)==final.read(name),name
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
audit={'source_preserved':True,'source_sha256':source_hash,'photos':20,'diagrams':10,
       'intended_pages':29,'output':str(OUTPUT),'footer_references':0}
(ART.parent/'document-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'SCREEN_PLAN_V12_CREATED: {OUTPUT} / all footers removed / 10 obstacle diagrams')

"""Screen specification with actual V10 player frames. Run with bundled DOCX Python."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[2]
IMAGES=ROOT/'docs/previews/ArtV10'
OUTPUT=ROOT/'docs/무한옥_화면별_게임기획서_V10.docx'
screens=[
('타이틀과 게임 시작','Title',
 '무한옥은 광산, 수레 철로, 계단, 감옥, 사육장을 달리는 운동 게임이다. 실제 화면을 기준으로 공간 구성과 조작, 애니메이션, 화면 전환을 정의한다.',
 [('화면 구성','나무 지지대와 주황색 랜턴으로 광산을 보여준다. 원숭이와 곡괭이, 로고가 서로 가리지 않도록 배치한다.'),('입력과 전환','시작하면 몸무게 선택 후 로딩 인트로를 거쳐 실제 광산 게임으로 바로 연결된다.'),('공통 화면 기준','게임은 16대9 비율을 유지한다. 다른 창 비율에서는 여백을 두어 로고와 HUD가 잘리지 않게 한다.')]),
('선택 몸무게 입력','WeightSetup',
 '사용자는 시작 전에 몸무게를 입력하거나 건너뛸 수 있다. 운동량을 보려는 사용자는 입력하고, 바로 게임만 하려는 사용자는 빈칸으로 시작한다.',
 [('입력 규칙','몸무게는 10~350 kg 범위의 숫자로 입력한다. 값은 이번 실행의 메모리에만 보관한다. 빈칸은 허용하며, 잘못된 값으로는 시작하지 않는다.'),('칼로리 표시','몸무게가 있을 때 인식된 실제 운동 시간으로 추정 소모량을 표시한다. 현재 계산은 고정 4 MET를 사용한다. 개인별 실제 에너지 소비를 측정하는 기능은 아니다.'),('시간 처리','키보드 테스트, 대기, 로딩, 일시정지 시간은 운동 시간으로 누적하지 않는다. 시작하면 인트로 역할의 로딩 화면으로 이동한다.')]),
('로딩 인트로','Loading',
 '제공된 원숭이와 경찰의 광산 추격 이미지를 로딩 화면으로 사용한다. 세계관과 목표를 한눈에 보여 주면서 실제 맵과 캐릭터를 준비한다.',
 [('화면 구성','광산 이미지의 원래 비율을 보존하고 아래쪽 진행 표시를 배치한다. 준비가 완료되면 이미지를 서서히 걷어 실제 광산을 드러낸다.'),('연출 연결','이 로딩 화면 자체가 인트로다. 완료 뒤 별도의 8초 영상을 추가로 재생하지 않는다. 첫 플레이 위치와 카메라를 미리 준비해 갑작스러운 이동을 줄인다.'),('아트 기준','실제 광산도 나무 보강재, 철로, 수레, 암석, 따뜻한 랜턴을 사용한다. 그림 속 조명과 색감을 게임의 공간 구성 기준으로 삼는다.')]),
('로딩에서 실제 광산으로 전환','LoadingReveal',
 '로딩 이미지가 사라지는 동안 같은 광산 분위기의 실제 3D 화면이 나타난다. 이미지 전환이 끝나면 달리기가 시작된다.',
 [('전환 방식','로딩 이미지의 불투명도를 약 0.85초 동안 낮춘다. 실제 원숭이의 시작 위치와 1인칭 카메라가 먼저 준비된다.'),('입력 처리','전환 중에는 이동과 운동 시간을 진행하지 않는다. 완료 뒤 좌우 이동, 점프, 숙이기, 하이니 입력을 받을 수 있다.'),('공간 연결','광산 바닥과 레일, 랜턴을 화면 중심의 진행 방향에 맞춘다. 수레 탑승 구간은 앞쪽 승강장으로 자연스럽게 이어진다.')]),
('광산 달리기','FirstLiveFrame',
 '광산에서 1인칭 달리기를 시작한다. 앞으로 자동 이동하며, 신체 동작으로 장애물을 통과하고 코인을 수집한다.',
 [('입력과 장애물','낮은 상자는 점프, 머리 위 파이프는 숙이기, 계단형 장애물은 하이니로 통과한다. 차단벽과 레이저는 안전한 차선으로 피한다.'),('실패 조건','필요한 동작을 하지 않고 단단한 장애물에 한 번 부딪히면 즉시 잡힘 상태가 된다. 잘못된 동작을 여러 번 허용하는 체력 방식은 사용하지 않는다.'),('아트와 다음 구간','통로 옆에 광석과 바위, 수레를 두고 중앙 진행 공간을 확보한다. 승강장에 도착하면 전용 수레 철로로 자동 탑승한다.')]),
('전용 철로의 수레 탑승','CartRide',
 '수레 구간은 걷는 길과 별개인 좁은 단선 철로다. 절벽 위 다리를 달리고, 카메라가 뒤로 물러나 수레와 원숭이를 보여 준다.',
 [('동작과 카메라','원숭이는 수레에 앉아 손잡이를 잡는 전용 동작을 재생한다. 1인칭 손은 숨기고, 뒤쪽 카메라가 수레의 회전과 몸 기울이기를 따라간다.'),('입력 규칙','왼쪽과 오른쪽 입력은 수레의 몸 기울이기로 사용한다. 수레에서는 점프와 하이니를 사용하지 않는다. 일반 상자나 파이프 장애물은 배치하지 않는다.'),('운동 흐름','다음 위험을 읽으면서 몸을 좌우로 기울이는 구간이다. 철로 두 청크를 지나면 하차하여 다시 달리기와 계단 운동으로 이어진다.')]),
('무너진 철로 회피','CartBank',
 '철로의 한쪽 레일과 다리 바닥이 끊어진 곳에서는 남아 있는 쪽으로 몸을 기울여 통과한다.',
 [('성공 조건','왼쪽 레일이 없으면 오른쪽으로, 오른쪽 레일이 없으면 왼쪽으로 기울인다. 현재 판정은 요구 방향으로 0.55 이상 기울인 상태에서 통과하는 것이다.'),('실패 조건','잘못된 방향으로 기울이거나 기울이지 않으면 수레가 떨어지고 게임이 종료된다. 앞을 막는 장애물을 피하는 방식과 구분한다.'),('위험의 시각 표현','깨진 다리와 절벽, 한쪽만 남은 레일을 멀리서 알아볼 수 있게 한다. 카메라 회전은 진행 방향을 유지하며 수레의 기울기를 약하게 반영한다.')]),
('수레 하차와 시설 진입','CartExit',
 '수레 철로 끝의 승강장에서 내린 뒤, 광산의 암벽과 목재가 점차 시설의 콘크리트와 배관으로 바뀐다.',
 [('연결 구조','광산 끝과 시설 입구 사이에 승강장, 끊기는 레일 끝, 공용 출입 프레임을 배치한다. 바닥 높이와 진행 방향은 이어진다.'),('카메라 전환','현재 화면 카메라 위치를 다음 카메라의 시작점으로 사용해 시점 전환 때 튀는 이동을 줄인다. 하차 후에는 1인칭 손이 다시 나타난다.'),('다음 동작','계단 구간에 들어가기 전에 걷기용 차선과 운동 장애물로 돌아온다. 계단에서는 무릎을 높이 드는 하이니 동작을 요구한다.')]),
('광산에서 시설 계단으로 상승','Map_04',
 '시설 계단은 광산과 감옥의 높이 차이를 연결한다. 배관, 난간, 경고등이 위쪽의 감옥 입구를 안내한다.',
 [('맵 구조','24 m 구간에서 높이가 3 m 상승한다. 바닥과 계단 앞코, 양쪽 난간을 같은 경로에 맞춰 연결한다. 시작과 끝의 경사는 완만하게 이어진다.'),('동작과 접지','계단형 운동 장애물은 하이니로 통과한다. 캐릭터 발은 경로의 바닥 높이와 경사에 맞춰 보정하며, 들어 올린 발은 바닥으로 강제로 끌어내리지 않는다.'),('다음 화면','계단 끝에서 감옥 입구의 문이 내려오기 시작한다. 사용자는 숙이기 안내를 보고 슬라이딩 동작으로 입구를 통과한다.')]),
('계단 끝에서 닫히는 감옥 문','ClosingGate',
 '감옥 입구를 향해 달릴 때 머리 위 문이 내려온다. 계단에서 복도로 넘어가는 순간에 숙이기 동작을 명확하게 요구한다.',
 [('문 움직임','문까지 14 m에서 3 m 사이의 접근 거리 동안 아래 모서리가 높은 위치에서 약 0.95 m까지 내려온다. 좌우 차선으로 돌아서 피할 수 없는 입구 문이다.'),('시각 안내','입구 프레임과 경고색 문 패널, 양쪽 상태등으로 닫히는 방향을 보여 준다. 계단 위 감옥의 차가운 조명이 다음 공간을 예고한다.'),('성공과 실패','숙인 채 문 아래를 지나가면 감옥으로 진입한다. 서서 닿거나 점프로 통과하려 하면 단단한 장애물 충돌로 종료된다.')]),
('문 아래 슬라이딩','SlideUnderGate',
 '원숭이가 몸을 낮추고 문 아래를 통과한다. 1인칭 카메라도 함께 낮아져 동작 결과를 체감할 수 있다.',
 [('애니메이션','실제 리타게팅된 숙이기 동작을 재생한다. 정상 달리기에서 짧게 블렌딩하고, 입력 시간이 끝나면 다시 달리기로 돌아온다.'),('공간의 연속성','입구 바로 다음에 감옥 복도가 시작한다. 통과한 뒤 3 m를 지나면 문은 바닥까지 닫혀 뒤쪽 광산과 계단을 막는다.'),('판정 기준','문을 통과하는 순간의 숙이기 상태로 성공을 판단한다. 화면 연출과 충돌 높이가 서로 다른 결과를 만들지 않도록 같은 위치에서 검사한다.')]),
('철창이 있는 감옥 복도','PrisonEntry',
 '감옥은 양쪽 독방과 중앙 통로가 구분되는 공간이다. 침대는 철창 안 깊은 위치에 두고, 원숭이는 잠긴 문 뒤에서 복도를 바라본다.',
 [('독방 구성','독방마다 바닥, 뒷벽, 양쪽 칸막이, 문틀과 철창을 만든다. 문에 경첩과 잠금장치, 수감 번호판을 배치해 닫힌 방으로 읽히게 한다.'),('공간 분리','복도 앞 철창은 중심에서 3.97 m, 뒷벽은 약 7.9 m에 있다. 침대와 담요, 벤치와 물그릇은 방 안에 놓고 달리기 통로를 침범하지 않는다.'),('동작과 다음 공간','원숭이는 철창 뒤에 머문다. 복도 중앙의 장애물과 차선별 닫히는 문을 통과하면 보안 구간과 사육장 출구로 이어진다.')]),
('철창 안 원숭이의 상태','CellCloseup',
 '원숭이를 철창 뒤에 배치하여 실제로 수감된 상태를 보여 준다. 철창 앞 복도에서 촬영한 공간 검수 화면이다.',
 [('캐릭터 동작','독방 원숭이는 철창 쪽으로 팔을 뻗은 채 작게 움직이고 고개를 좌우로 돌린다. 얼굴의 눈 깜빡임과 입 움직임도 유지한다.'),('배치 기준','몸과 발이 철창 바깥으로 나오지 않게 배치한다. 침대는 뒤쪽 벽, 물그릇은 방 안 바닥에 놓는다. 원숭이가 중앙 통로의 충돌 판정에 끼어들지 않는다.'),('분위기','차가운 복도 조명과 따뜻한 독방 랜턴을 함께 사용해 방의 깊이를 드러낸다. 이 근접 카메라는 제작 검수용이며 일반 플레이는 1인칭이다.')]),
('캐릭터 디테일과 발 접지','CharacterContact',
 '실제 광산 안에서 원숭이를 옆에서 촬영한 캐릭터 검수 화면이다. 달리기, 발의 높이, 손발과 얼굴의 디테일을 확인한다.',
 [('외형','갈색 털과 밝은 얼굴·손바닥·발바닥을 구분한다. 손목과 허리에 주황색 수감 밴드, 등에 0723 번호 패치를 추가했다. 손발톱, 입과 이빨, 꼬리와 눈 표정을 갖춘 스킨 모델이다.'),('접지 동작','낮은 발은 바닥 높이에 맞추고 발바닥 회전은 경사에 맞춘다. 접지 순간에는 짧은 범위 안에서 발 위치를 유지해 미끄러짐을 줄인다. 큰 이동에서는 고정을 풀어 다리가 늘어나지 않게 한다.'),('동작 구분','점프와 수레에서는 바닥 접지 보정을 해제한다. 달리기 속도는 감속 상태에 맞춰 낮춘다. 몸 이동은 게임 규칙이 담당하고 애니메이션은 뼈대에 적용한다.')]),
('감옥 보안 구간과 출구','Map_07',
 '감옥의 끝은 사육장으로 나가는 관리 시설이다. 검사 장치와 출입 프레임으로 복도에서 외부 공간으로 넘어가는 이유를 만든다.',
 [('공간 구성','양쪽에 검사 장치와 상태등을 배치하고 중앙 통로는 확보한다. 콘크리트 벽, 배관, 출입 프레임을 뒤쪽 외부 출구까지 이어 놓는다.'),('입력과 장애물','문 패널이 내려오는 차선은 옆으로 이동한다. 전 구간 숙이기 장애물은 해당 동작으로 통과한다. 잘못된 동작으로 단단한 장애물에 닿으면 종료된다.'),('다음 맵 연결','관리 출구를 통과하면 사육장의 완만한 내리막으로 연결한다. 바닥 높이가 3 m에서 0 m로 낮아지며 야외의 달빛과 수목이 보이기 시작한다.')]),
('밤의 동물 사육장','Map_09',
 '철제 울타리와 동물 우리 사이의 야외 길을 달린다. 감옥의 폐쇄된 복도에서 밤하늘이 보이는 사육장으로 공간이 넓어진다.',
 [('아트 구성','양쪽 울타리 뒤에 기린과 코끼리, 원숭이 공간을 두고 나무와 식물을 배치한다. 따뜻한 랜턴과 푸른 밤하늘이 길의 방향을 보여 준다.'),('장애물과 바닥 아이템','통나무는 점프, 바나나와 물웅덩이는 피하거나 뛰어넘는다. 단단한 장애물과 감속 아이템이 서로 다른 결과를 만든다.'),('반복 경로','출구를 지나면 다음 광산 경로로 이어지는 반복 맵이 계속된다. 이전 구간과 다음 구간의 높이, 폭, 연결점은 같은 규칙을 사용한다.')]),
('바나나를 밟은 첫 감속','Slip',
 '바나나를 밟으면 바로 종료하지 않고 잠시 느려진다. 경찰이 가까워지는 위험을 HUD와 비틀거림으로 전달한다.',
 [('첫 접촉','현재 감속은 2.4초 동안 속도 50%로 적용된다. 바나나 접촉 횟수와 추격 압력이 증가하며, 감속이 끝나도 접촉 횟수는 이번 플레이에 남는다.'),('회피 방법','옆 차선으로 이동하거나 점프하여 바닥 아이템을 피한다. 물웅덩이도 감속하지만 바나나 접촉 횟수에는 포함되지 않는다.'),('애니메이션','비틀거림 상태에서 얼굴과 몸의 반응을 보여 주고, 감속 중 달리기 동작의 속도를 낮춘다. 정상 속도로 회복하면 다시 기존 달리기로 연결한다.')]),
('바나나 두 번 접촉 후 경찰에게 잡힘','BananaCatch',
 '같은 플레이에서 바나나를 두 번 이상 밟으면 감속으로 경찰이 따라잡는다. 바나나 접촉과 추격의 결과가 연결되어야 한다.',
 [('규칙','두 번째 접촉 뒤 감속과 추격 연출이 진행되고 약 1.2초 뒤 잡힘 상태가 된다. 첫 감속이 이미 끝났어도 같은 플레이의 두 번째 접촉이면 적용한다.'),('잡힘 화면','경찰이 원숭이 앞쪽으로 다가와 체포 동작을 재생한다. 카메라 안으로 얼굴이 겹쳐 들어오지 않도록 간격을 확보한다.'),('재시작','짧은 잡힘 연출 후 결과 화면으로 전환한다. 다시 시작하면 바나나 접촉 횟수, 감속, 추격, 획득 점수와 현재 운동 기록을 초기화한다.')]),
('일시정지와 재개','Pause',
 '일시정지에서는 게임 상태와 운동 시간을 멈춘다. 사용자가 다시 시작하면 같은 위치에서 이어 달린다.',
 [('화면 구성','현재 맵 위에 일시정지 표시와 재개 버튼을 보여 준다. 뒤쪽 실제 게임 화면을 남겨 현재 위치와 다음 장애물을 알아볼 수 있게 한다.'),('정지 범위','자동 전진, 추격, 감속 시간과 운동 시간은 함께 멈춘다. 정지 전에 쌓인 입력은 지워 재개 직후 예상하지 못한 동작이 나오는 것을 막는다.'),('재개 조건','재개 버튼이나 대응 키로 이어 한다. 멈춰 있는 동안 운동 입력만으로 칼로리 추정치가 계속 올라가지 않는다.')]),
('결과 화면과 다시 시작','GameOver',
 '단단한 장애물 충돌, 수레 추락, 바나나 감속 추격의 종료 결과를 보여 준다. 사용자는 기록을 확인하고 바로 다시 시작할 수 있다.',
 [('표시 내용','거리와 점수, 최고 기록을 표시한다. 몸무게를 입력한 경우에는 인식된 운동 시간을 바탕으로 계산한 추정 칼로리 정보를 확인한다.'),('종료 구분','단단한 장애물은 한 번 충돌하면 종료된다. 수레는 잘못된 기울이기로 철로 붕괴 지점을 통과했을 때 추락한다. 바나나는 두 번째 접촉 후 추격으로 종료된다.'),('다시 시작','재시작하면 타이틀로 돌아가 몸무게 선택과 로딩을 거친다. 장애물, 코인, 바닥 아이템과 문 상태도 다시 준비한다.')]),
]

doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.65);section.bottom_margin=Inches(.6)
section.left_margin=section.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style=doc.styles[name];style.font.name='맑은 고딕';style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():del fonts.attrib[attr]
    for tag in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+tag),'Malgun Gothic')
    style.font.italic=False;style.font.bold=name in ['Heading 1','Heading 2']
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.20;normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(22)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.core_properties.title='무한옥 화면별 게임 기획서'
doc.core_properties.subject='캐릭터 디테일 발 접지 독방 공간 게임 화면과 동작 흐름'
doc.core_properties.author='무한옥 프로젝트'
doc.core_properties.comments=''

footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('무한옥 V10   ·   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.style='Caption'

for index,(title,image,intro,blocks) in enumerate(screens,1):
    p=doc.add_paragraph('무한옥 화면별 게임 기획서' if index==1 else f'{index:02d} {title}', 'Title' if index==1 else 'Heading 1')
    if index>1:p.paragraph_format.page_break_before=True
    if index==1:
        doc.add_paragraph('실제 게임 화면과 캐릭터 공간 동작 설계','Subtitle')
        meta=doc.add_paragraph('2026년 10월 5일   |   Unity V10   |   화면 01 타이틀', 'Caption')
    else:doc.add_paragraph(f'화면 {index:02d}   |   실제 Unity 실행 화면', 'Caption')
    path=IMAGES/f'Player_{image}_1280x720.png';assert path.exists(),path
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    if index==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(6.7 if index==1 else 7.1))
    picture=p.runs[0]._r.xpath('.//wp:docPr')[0];picture.set('descr',f'{title} 실제 Unity 화면')
    p.paragraph_format.keep_with_next=True
    label='제작 검수용 근접 카메라' if image in ['CellCloseup','CharacterContact'] else '실제 플레이 화면'
    doc.add_paragraph(f'그림 {index:02d} {title}   |   {label}', 'Caption')
    doc.add_paragraph(intro)
    for heading,body in blocks:
        doc.add_paragraph(heading,'Heading 2');doc.add_paragraph(body)

p=doc.add_paragraph('21 동작별 장애물과 판정','Heading 1');p.paragraph_format.page_break_before=True
doc.add_paragraph('사용자가 미리 읽을 수 있는 형태와 실제 동작 판정을 연결한다. 같은 장애물이 구간에 따라 다른 동작을 요구하지 않도록 한다.')
table=doc.add_table(rows=1,cols=3);table.autofit=False
for cell,text,width in zip(table.rows[0].cells,['대상','필요한 동작','통과하지 못했을 때'],[1.55,2.45,3.1]):
    cell.width=Inches(width);cell.text=text
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7ECF1');cell._tc.get_or_add_tcPr().append(fill)
    for run in cell.paragraphs[0].runs:run.bold=True
for row in [('낮은 상자와 통나무','점프하여 넘기','한 번 충돌하면 종료'),('머리 위 파이프','몸을 낮추기','한 번 충돌하면 종료'),('계단형 운동 장애물','무릎 높이 들기','한 번 충돌하면 종료'),('감옥 입구의 닫히는 문','숙여서 아래로 통과','서서 닿으면 종료'),('차선 차단벽과 레이저','안전한 차선으로 이동','한 번 충돌하면 종료'),('수레의 붕괴 철로','남은 레일 쪽으로 기울이기','수레 추락으로 종료'),('바나나','피하거나 점프로 넘기','첫 접촉 감속 두 번째 추격 종료'),('물웅덩이','피하거나 점프로 넘기','일시 감속 바나나 횟수는 유지')]:
    for cell,text in zip(table.add_row().cells,row):cell.text=text
borders=OxmlElement('w:tblBorders')
for tag in ['top','left','bottom','right','insideH','insideV']:
    b=OxmlElement('w:'+tag);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
table._tbl.tblPr.append(borders)
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
doc.add_paragraph('개발용 키보드 조작','Heading 2')
doc.add_paragraph('좌우 방향키 또는 A와 D로 이동한다. Space로 점프, 아래 방향키 또는 S로 숙이기, 위 방향키 또는 W로 하이니를 테스트한다. 실제 운동은 포즈 입력을 통해 같은 게임 동작으로 연결한다.')
doc.add_paragraph('공간과 애니메이션 확인','Heading 2')
doc.add_paragraph('광산에서 수레 철로, 계단, 닫히는 문, 감옥, 사육장으로 이어지는 경계를 확인한다. 독방 원숭이의 발은 바닥에 있고 몸은 철창 뒤에 있어야 한다. 점프와 수레 동작에서는 발 접지 보정이 해제되어야 한다.')
OUTPUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUTPUT)
print(f'SCREEN_PLAN_CREATED: {OUTPUT} / {len(screens)} screen photos / 21 intended pages')
